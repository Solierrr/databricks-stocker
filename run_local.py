"""Executa o medalhão em Spark local e exporta os KPIs Gold para JSON."""

import argparse
from collections import Counter
from contextlib import ExitStack
from datetime import date, datetime, timezone
from decimal import Decimal
import importlib
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
import types
from unittest.mock import patch

ROOT = Path(__file__).resolve().parent
TRANSFORMATIONS = ROOT / "pipelines/00_medallion/transformations"


def json_value(value):
    if isinstance(value, Decimal):
        # Texto preserva a precisão dos valores monetários.
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    raise TypeError(f"Tipo não suportado no JSON: {type(value).__name__}")


class LocalReader:
    """Adapta somente a leitura de Excel; os tratamentos Bronze continuam originais."""

    def __init__(self, session):
        self.session = session
        self.options = {}

    def table(self, name):
        return self.session.read.table(name)

    def format(self, name):
        if name != "excel":
            raise ValueError(f"Formato local não implementado: {name}")
        return self

    def option(self, name, value):
        self.options[name] = value
        return self

    def load(self, path):
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            name = self.options["dataAddress"]
            if name not in book:
                raise ValueError(f"{Path(path).name}: aba ausente: {name}")
            rows = book[name].iter_rows(values_only=True)
            header = next(rows, ())
            if not header:
                raise ValueError(f"{name}: cabeçalho ausente")
            positions = [i for i, v in enumerate(header) if v is not None and str(v).strip()]
            names = [str(header[i]) for i in positions]
            if len(set(names)) != len(names):
                raise ValueError(f"{name}: cabeçalhos duplicados")
            values = []
            for row in rows:
                if any(v is not None for v in row):
                    values.append(tuple(self._text(row[i] if i < len(row) else None) for i in positions))
            # VALUES evita depender de processos Python auxiliares no Windows.
            # Literais hexadecimais preservam acentos, aspas e barras com segurança.
            columns = ", ".join("`" + n.replace("`", "``") + "`" for n in names)
            if values:
                encoded = ", ".join("(" + ", ".join(self._literal(v) for v in row) + ")" for row in values)
                decoded = ", ".join(f"decode(unhex(`{n.replace('`', '``')}`), 'UTF-8') AS `{n.replace('`', '``')}`" for n in names)
                frame = self.session.sql(f"SELECT {decoded} FROM VALUES {encoded} AS local_excel({columns})")
            else:
                fields = ", ".join(f"CAST(NULL AS STRING) AS `{n.replace('`', '``')}`" for n in names)
                frame = self.session.sql(f"SELECT {fields} WHERE FALSE")
            frame = frame.cache()
            if Counter(tuple(row) for row in frame.collect()) != Counter(values):
                frame.unpersist()
                raise RuntimeError(f"{name}: a leitura local divergiu do Excel; exportação cancelada")
            return frame
        finally:
            book.close()

    @staticmethod
    def _literal(value):
        if value is None:
            return "CAST(NULL AS STRING)"
        return "'" + value.encode("utf-8").hex() + "'"

    @staticmethod
    def _text(value):
        if value is None:
            return None
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        if isinstance(value, bool):
            return str(value).lower()
        return str(value)


class LocalSpark:
    def __init__(self, session):
        self.session = session
        self.tables = set()
        self.catalog = types.SimpleNamespace(tableExists=lambda name: name in self.tables)

    @property
    def read(self):
        return LocalReader(self.session)

    def sql(self, query):
        # O executor usa apenas views temporárias, sem schemas persistentes.
        if query.startswith("CREATE SCHEMA IF NOT EXISTS "):
            return None
        raise ValueError(f"Operação persistente bloqueada no executor local: {query[:80]}")


class LocalWidgets:
    def __init__(self, parameters):
        self.parameters = parameters

    def text(self, name, default):
        self.parameters.setdefault(name, default)

    def get(self, name):
        return self.parameters[name]


def run(args):
    # A configuração vale apenas para este processo e para o Java iniciado por ele.
    if args.temp_dir:
        temp = Path(args.temp_dir).resolve()
        temp.mkdir(parents=True, exist_ok=True)
        os.environ["TEMP"] = os.environ["TMP"] = str(temp)
        os.environ["JAVA_TOOL_OPTIONS"] = os.environ.get("JAVA_TOOL_OPTIONS", "") + f' -Djava.io.tmpdir="{temp}"'
    for path in (args.fornecedores, args.admins):
        if not path.is_file():
            raise FileNotFoundError(f"Excel não encontrado: {path}")
    if not shutil.which("java"):
        raise RuntimeError("Instale Java 17 ou 21 e disponibilize java no PATH.")
    from pyspark.sql import SparkSession

    os.environ.setdefault("PYSPARK_PYTHON", sys.executable)
    session = (SparkSession.builder.master("local[2]").appName("solaria-local")
               .config("spark.ui.enabled", "false")
               .config("spark.sql.shuffle.partitions", "2")
               .config("spark.sql.session.timeZone", "America/Sao_Paulo")
               .config("spark.sql.codegen.factoryMode", "NO_CODEGEN")
               .config("spark.sql.adaptive.enabled", "false")
               .getOrCreate())
    session.sparkContext.setLogLevel("ERROR")
    runtime = types.ModuleType("databricks.sdk.runtime")
    runtime.spark = LocalSpark(session)
    runtime.dbutils = types.SimpleNamespace(widgets=LocalWidgets({
        "catalog": "local", "source_mode": "manual",
        "fornecedores_path": str(args.fornecedores), "admins_path": str(args.admins),
        "allow_replace_existing": "false",
    }))
    result = {"metadata": {"generated_at": datetime.now(timezone.utc).isoformat(),
                           "timezone": "America/Sao_Paulo", "mode": "local",
                           "decimals": "strings", "supplier_access_filter_applied": False,
                           "inputs": {"fornecedores": str(args.fornecedores), "admins": str(args.admins)}},
              "tables": {}}
    cached = []
    try:
        with ExitStack() as stack:
            stack.enter_context(patch.dict(sys.modules, {"databricks.sdk.runtime": runtime}))
            stack.enter_context(patch.object(sys, "path", [str(TRANSFORMATIONS), *sys.path]))
            catalog = importlib.import_module("note_catalog")
            stack.enter_context(patch.object(catalog, "table", lambda schema, name: f"local_{schema}_{name}"))
            main = importlib.import_module("my_transformation")

            def write_local(name, frame, *, supplier_filter=False):
                if name.startswith("local_gold_"):
                    rows = [row.asDict(recursive=True) for row in frame.collect()]
                    # As tabelas não têm ordem garantida; padronizamos só a saída do JSON.
                    rows.sort(key=lambda row: json.dumps(row, sort_keys=True, default=json_value))
                    key = name.removeprefix("local_gold_")
                    result["tables"][key] = rows
                    print(f"Gold: {key} ({len(rows)} linhas)", flush=True)
                else:
                    materialized = frame.cache()
                    cached.append(materialized)
                    count = materialized.count()
                    materialized.createOrReplaceTempView(name)
                    runtime.spark.tables.add(name)
                    print(f"{name}: {count} linhas", flush=True)

            stack.enter_context(patch.object(main, "_write_table", write_local))
            main.run()
        # Publica o JSON somente se o fluxo completo terminar com sucesso.
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(mode="w", encoding="utf-8", dir=args.output.parent,
                                         suffix=".tmp", delete=False) as handle:
            temporary = Path(handle.name)
            try:
                json.dump(result, handle, ensure_ascii=False, indent=2, default=json_value, allow_nan=False)
            except BaseException:
                handle.close()
                temporary.unlink(missing_ok=True)
                raise
        temporary.replace(args.output)
        print(f"\nJSON salvo em: {args.output}", flush=True)
    finally:
        session.stop()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fornecedores", type=Path, default=ROOT / "recovery/fornecedores.xlsx")
    parser.add_argument("--admins", type=Path, default=ROOT / "recovery/admins.xlsx")
    parser.add_argument("--output", type=Path, default=ROOT / "local-output/gold_kpis.json")
    parser.add_argument("--temp-dir", help="Pasta temporária curta para Java no Windows, se necessário")
    args = parser.parse_args()
    args.fornecedores = args.fornecedores.resolve()
    args.admins = args.admins.resolve()
    args.output = args.output.resolve()
    try:
        run(args)
    except ImportError as error:
        print(f"Dependência ausente: {error}. Instale: python -m pip install -r requirements-local.txt", file=sys.stderr)
        return 1
    except Exception as error:
        print(f"Execução local falhou: {error}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
